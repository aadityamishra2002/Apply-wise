from io import BytesIO
from types import SimpleNamespace

from docx import Document
from docx.shared import Pt

from app.application.exports.service import ResumeExportService, extract_placeholders_from_docx, infer_section_order_from_docx


def make_resume_template_docx() -> bytes:
    document = Document()
    document.add_paragraph().add_run("{{NAME}}").font.size = Pt(18)
    document.add_paragraph().add_run("{{CONTACT}}")
    document.add_heading("Profile", level=1)
    summary = document.add_paragraph()
    summary_run = summary.add_run("{{SUMMARY}}")
    summary_run.bold = True
    document.add_heading("Selected Experience", level=1)
    document.add_paragraph(style="List Bullet").add_run("{{EXPERIENCE}}")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_plain_resume_template_docx() -> bytes:
    document = Document()
    document.add_heading("Summary", level=1)
    document.add_paragraph("Old summary text that should be replaced.")
    document.add_heading("Experience", level=1)
    document.add_paragraph("Old role", style="List Bullet")
    document.add_paragraph("Old experience text", style="List Bullet")
    document.add_heading("Education", level=1)
    document.add_paragraph("Old education line")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_template_renderer_replaces_placeholders_and_preserves_paragraph_style() -> None:
    template_bytes = make_resume_template_docx()
    assert extract_placeholders_from_docx(template_bytes) == ["NAME", "CONTACT", "SUMMARY", "EXPERIENCE"]
    service = ResumeExportService()
    version = SimpleNamespace(
        resume_content={
            "professional_summary": "Product engineer building reliable user-facing software.",
            "skills": ["Python", "React"],
            "experience": [{"title": "Senior Product Engineer", "company": "Northstar Labs", "description": "Built FastAPI services."}],
            "projects": [{"name": "Apply Wise Copilot", "description": "Built resume workflows.", "technologies": ["Python"]}],
            "education": [{"institution": "State University", "degree": "Bachelor of Engineering", "year": "2021"}],
            "certifications": [{"name": "AWS Certified Cloud Practitioner", "issuer": "Amazon Web Services", "year": "2022"}],
        }
    )
    job = SimpleNamespace(title="Platform Engineer", company="Example")
    user = SimpleNamespace(email="owner@example.com")
    profile = SimpleNamespace(full_name="Aaditya Mishra", phone="", location="India", linkedin_url="", github_url="", portfolio_url="")
    template = SimpleNamespace(template_bytes=template_bytes, placeholders=["NAME", "CONTACT", "SUMMARY", "EXPERIENCE"])

    output = service.build_docx(version, job, user=user, user_profile=profile, template=template)
    document = Document(BytesIO(output))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert "{{" not in text
    assert "Aaditya Mishra" in text
    assert "Northstar Labs" in text
    assert "Apply Wise Copilot" in text
    assert "AWS Certified Cloud Practitioner" in text
    summary = next(paragraph for paragraph in document.paragraphs if "Product engineer" in paragraph.text)
    assert summary.runs[0].bold is True
    assert any(paragraph.style.name == "List Bullet" and "Senior Product Engineer" in paragraph.text for paragraph in document.paragraphs)


def test_plain_docx_template_replaces_existing_sections_without_rebuilding_layout() -> None:
    template_bytes = make_plain_resume_template_docx()
    assert extract_placeholders_from_docx(template_bytes) == []
    assert infer_section_order_from_docx(template_bytes) == ["SUMMARY", "EXPERIENCE", "EDUCATION"]
    service = ResumeExportService()
    version = SimpleNamespace(
        resume_content={
            "professional_summary": "Product engineer building reliable user-facing software.",
            "skills": ["Python", "React"],
            "experience": [{"title": "Senior Product Engineer", "company": "Northstar Labs", "description": "Built FastAPI services."}],
            "projects": [],
            "education": [{"institution": "State University", "degree": "Bachelor of Engineering", "year": "2021"}],
            "certifications": [],
        }
    )
    job = SimpleNamespace(title="Platform Engineer", company="Example")
    template = SimpleNamespace(template_bytes=template_bytes, placeholders=[])

    output = service.build_docx(version, job, template=template)
    document = Document(BytesIO(output))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert "Old summary text" not in text
    assert "Product engineer building reliable" in text
    assert "Northstar Labs" in text
    assert "State University" in text
    assert "Summary" in text
    assert not text.startswith("Aaditya")
    assert any(paragraph.style.name == "List Bullet" and "Senior Product Engineer" in paragraph.text for paragraph in document.paragraphs)


def test_default_template_export_renders_all_supported_sections() -> None:
    service = ResumeExportService()
    version = SimpleNamespace(
        resume_content={
            "professional_summary": "Backend engineer focused on product reliability.",
            "skills": ["Python"],
            "experience": [{"title": "Engineer", "company": "Northstar Labs", "description": "Built APIs."}],
            "projects": [{"name": "Apply Wise", "description": "Built templates.", "technologies": []}],
            "education": [{"institution": "State University", "degree": "BE", "year": "2021"}],
            "certifications": [{"name": "AWS CCP", "issuer": "AWS", "year": "2022"}],
        }
    )
    job = SimpleNamespace(title="Engineer", company="Example")

    output = service.build_docx(version, job)
    document = Document(BytesIO(output))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert "Backend engineer focused" in text
    assert "Northstar Labs" in text
    assert "Apply Wise" in text
    assert "State University" in text
    assert "AWS CCP" in text
