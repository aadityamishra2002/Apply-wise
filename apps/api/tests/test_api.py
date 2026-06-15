import asyncio
from datetime import datetime, timedelta, timezone
from io import BytesIO
from uuid import UUID, uuid4

import fitz
import httpx
import jwt
import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy.exc import SQLAlchemyError

from app.core.config import get_settings
from app.infrastructure.db.models import Job, JobMatch
from app.infrastructure.db.session import session_factory
from app.main import app


client = TestClient(app)


def test_health() -> None:
    response = client.get("/api/v1/health")
    assert response.status_code == 200
    assert response.json()["status"] == "ok"
    assert response.headers["x-request-id"]


def test_readiness_checks_database() -> None:
    response = client.get("/api/v1/health/ready")
    assert response.status_code == 200
    assert response.json() == {"status": "ready", "database": "ok"}


def test_feature_health_does_not_expose_credentials() -> None:
    response = client.get("/api/v1/health/features")
    assert response.status_code == 200
    assert response.json() == {"adzuna_enabled": False, "gemini_enabled": False, "remotive_enabled": False}


def test_provider_health_is_public_and_non_secret() -> None:
    response = client.get("/api/v1/health/providers")
    assert response.status_code == 200
    assert response.json()[0]["name"] == "mock"
    assert response.json()[0]["status"] in {"configured", "available"}
    assert "app_key" not in response.text


def test_automation_capabilities_are_public_and_manual_submit_only() -> None:
    response = client.get("/api/v1/automation/capabilities")
    assert response.status_code == 200
    capabilities = {item["site_key"]: item for item in response.json()}
    assert capabilities["linkedin"]["capture_supported"] is True
    assert capabilities["linkedin"]["dry_run_supported"] is True
    assert capabilities["linkedin"]["manual_submit_required"] is True
    assert capabilities["workday"]["autofill_supported"] is True
    assert capabilities["jobvite"]["autofill_supported"] is True
    assert capabilities["icims"]["manual_submit_required"] is True
    assert capabilities["bamboohr"]["capture_supported"] is True
    assert capabilities["workable"]["dry_run_supported"] is True
    assert capabilities["generic"]["autofill_supported"] is False


def test_local_browser_origin_is_allowed() -> None:
    response = client.options(
        "/api/v1/auth/register",
        headers={
            "Origin": "http://127.0.0.1:3000",
            "Access-Control-Request-Method": "POST",
            "Access-Control-Request-Headers": "content-type",
        },
    )
    assert response.status_code == 200
    assert response.headers["access-control-allow-origin"] == "http://127.0.0.1:3000"


def test_jobs_are_ranked() -> None:
    token = register_user()["tokens"]["access_token"]
    upload_resume(token)
    response = client.get("/api/v1/jobs", headers={"Authorization": f"Bearer {token}"})
    assert response.status_code == 200
    scores = [job["match_score"] for job in response.json()]
    assert scores == sorted(scores, reverse=True)
    assert response.json()[0]["ats_breakdown"]["keywords"] >= 0
    assert isinstance(response.json()[0]["high_fit"], bool)
    assert response.json()[0]["recommendation_summary"]
    assert any("resume skills" in reason for reason in response.json()[0]["recommendation_reasons"])
    summary = client.get("/api/v1/dashboard/summary", headers={"Authorization": f"Bearer {token}"}).json()
    assert summary["average_match_score"] is not None
    activity = client.get("/api/v1/activity/logs", headers={"Authorization": f"Bearer {token}"}).json()
    assert activity[0]["event_type"] == "job_search_scored"
    assert activity[0]["metadata"]["result_count"] == len(response.json())


def test_dashboard_summary_starts_empty() -> None:
    token = register_user()["tokens"]["access_token"]
    response = client.get("/api/v1/dashboard/summary", headers={"Authorization": f"Bearer {token}"})
    assert response.status_code == 200
    assert response.json() == {
        "applications": 0,
        "interviews": 0,
        "saved_jobs": 0,
        "average_match_score": None,
        "weekly_activity": [0, 0, 0, 0, 0, 0, 0],
        "has_resume": False,
        "has_application_profile": False,
        "prepared_applications": 0,
        "reviewed_applications": 0,
    }


def test_account_export_includes_workspace_data_without_auth_secrets() -> None:
    token = register_user("export-owner@example.com")["tokens"]["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    upload_resume(token)
    jobs = client.get("/api/v1/jobs", headers=headers)
    assert jobs.status_code == 200
    saved = client.post(f"/api/v1/jobs/{jobs.json()[0]['id']}/save", headers=headers)
    assert saved.status_code == 201
    exported = client.get("/api/v1/account/export", headers=headers)
    assert exported.status_code == 200
    assert exported.headers["content-type"].startswith("application/json")
    body = exported.json()
    assert body["account"]["email"] == "export-owner@example.com"
    assert body["resume_profile"]["skills"]
    assert body["applications"][0]["job"]["title"]
    assert body["job_matches"]
    serialized = exported.text.lower()
    assert "password_hash" not in serialized
    assert "refresh_token" not in serialized
    assert "token_hash" not in serialized


def test_analytics_summary_starts_empty() -> None:
    token = register_user()["tokens"]["access_token"]
    response = client.get("/api/v1/analytics/summary", headers={"Authorization": f"Bearer {token}"})
    assert response.status_code == 200
    assert response.json()["total_applications"] == 0
    assert response.json()["interview_rate"] == 0
    assert response.json()["offer_rate"] == 0


def test_job_preferences_are_persisted_and_used_for_default_discovery() -> None:
    token = register_user()["tokens"]["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    upload_resume(token)
    empty = client.get("/api/v1/settings/job-preferences", headers=headers)
    assert empty.status_code == 200
    assert empty.json() == {"role_keywords": "", "location": "", "remote_only": False}

    saved = client.put(
        "/api/v1/settings/job-preferences",
        headers=headers,
        json={"role_keywords": "backend engineer", "location": "remote", "remote_only": True},
    )
    assert saved.status_code == 200
    assert saved.json()["role_keywords"] == "backend engineer"
    jobs = client.get("/api/v1/jobs", headers=headers)
    assert jobs.status_code == 200
    assert [job["title"] for job in jobs.json()] == ["Backend Engineer"]
    overridden = client.get("/api/v1/jobs?query=engineer&location=Remote", headers=headers)
    assert overridden.status_code == 200
    assert overridden.json()
    assert all("remote" in job["location"].lower() for job in overridden.json())


def test_application_profile_is_persisted_exported_and_used_for_review_autofill() -> None:
    token = register_user("profile-autofill@example.com")["tokens"]["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    upload_resume(token)
    empty = client.get("/api/v1/settings/profile", headers=headers)
    assert empty.status_code == 200
    assert empty.json()["full_name"] == ""

    saved_profile = client.put(
        "/api/v1/settings/profile",
        headers=headers,
        json={
            "full_name": "Aaditya Mishra",
            "phone": "+91 98765 43210",
            "location": "Bengaluru, India",
            "linkedin_url": "https://www.linkedin.com/in/aaditya-mishra",
            "github_url": "https://github.com/aadityamishra2002",
            "portfolio_url": "https://applywise.app",
        },
    )
    assert saved_profile.status_code == 200
    assert saved_profile.json()["phone"] == "+91 98765 43210"

    job = client.get("/api/v1/jobs", headers=headers).json()[0]
    application = client.post(f"/api/v1/jobs/{job['id']}/save", headers=headers).json()
    review = client.get(f"/api/v1/applications/{application['id']}/review-before-apply", headers=headers)
    assert review.status_code == 200
    fields = {field["field_key"]: field for field in review.json()["autofill_fields"]}
    assert fields["full_name"]["value"] == "Aaditya Mishra"
    assert fields["phone"]["value"] == "+91 98765 43210"
    assert fields["linkedin"]["value"] == "https://www.linkedin.com/in/aaditya-mishra"

    export = client.get("/api/v1/account/export", headers=headers)
    assert export.status_code == 200
    assert export.json()["application_profile"]["github_url"] == "https://github.com/aadityamishra2002"


def test_application_profile_rejects_partial_urls() -> None:
    token = register_user("profile-url-validation@example.com")["tokens"]["access_token"]
    response = client.put(
        "/api/v1/settings/profile",
        headers={"Authorization": f"Bearer {token}"},
        json={
            "full_name": "Aaditya Mishra",
            "phone": "",
            "location": "India",
            "linkedin_url": "linkedin.com/in/aaditya",
            "github_url": "",
            "portfolio_url": "",
        },
    )
    assert response.status_code == 422
    assert "full URL" in response.text


def test_mock_job_ids_are_stable() -> None:
    token = register_user()["tokens"]["access_token"]
    upload_resume(token)
    headers = {"Authorization": f"Bearer {token}"}
    first_response = client.get("/api/v1/jobs", headers=headers)
    second_response = client.get("/api/v1/jobs", headers=headers)
    first_ids = [job["id"] for job in first_response.json()]
    second_ids = [job["id"] for job in second_response.json()]
    assert first_ids == second_ids


def test_resume_analysis_detects_skills() -> None:
    response = client.post(
        "/api/v1/resumes/analyze",
        json={"text": "Summary Experience Skills Python React Docker Education"},
    )
    assert response.status_code == 200
    assert response.json()["skills"] == ["python", "react", "docker"]


def test_resume_upload_requires_authentication() -> None:
    response = client.post(
        "/api/v1/resumes/upload",
        files={"file": ("resume.docx", make_resume_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 401


def test_docx_resume_upload_extracts_and_persists_profile() -> None:
    token = register_user()["tokens"]["access_token"]
    upload = client.post(
        "/api/v1/resumes/upload",
        headers={"Authorization": f"Bearer {token}"},
        files={"file": ("resume.docx", make_resume_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert upload.status_code == 201
    assert upload.json()["original_filename"] == "resume.docx"
    assert upload.json()["parse_status"] == "review_required"
    assert upload.json()["detected_sections"] == ["summary", "experience", "education", "skills", "projects", "certifications"]
    assert upload.json()["skills"] == ["AWS", "Docker", "FastAPI", "PostgreSQL", "Python", "React", "TypeScript"]
    assert upload.json()["experiences"] == [
        {
            "title": "Senior Product Engineer",
            "company": "Northstar Labs",
            "description": "Built React and TypeScript applications with Python, FastAPI, PostgreSQL, and Docker.",
        }
    ]
    assert upload.json()["education"] == [
        {
            "institution": "State University",
            "degree": "Bachelor of Engineering",
            "year": "2021",
        }
    ]
    assert upload.json()["certifications"] == [
        {
            "name": "AWS Certified Cloud Practitioner",
            "issuer": "Amazon Web Services",
            "year": "2022",
        }
    ]
    assert upload.json()["projects"] == [
        {
            "name": "Apply Wise Copilot",
            "description": "Built a React and FastAPI workflow that parses resumes, scores jobs, and tracks applications.",
            "technologies": ["FastAPI", "React"],
        }
    ]
    activity = client.get("/api/v1/activity/logs", headers={"Authorization": f"Bearer {token}"}).json()
    event_types = [event["event_type"] for event in activity]
    assert {"resume_parsed", "resume_uploaded"}.issubset(set(event_types))
    parsed_event = next(event for event in activity if event["event_type"] == "resume_parsed")
    assert parsed_event["metadata"]["skill_count"] == 7
    assert parsed_event["metadata"]["project_count"] == 1

    profile = client.get(
        "/api/v1/resumes/me",
        headers={"Authorization": f"Bearer {token}"},
    )
    assert profile.status_code == 200
    assert profile.json() == upload.json()


def test_resume_profile_can_be_edited_and_reparsed() -> None:
    token = register_user("resume-edit@example.com")["tokens"]["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    upload_resume(token)

    edited = client.patch(
        "/api/v1/resumes/me",
        headers=headers,
        json={
            "skills": ["Python", "FastAPI", "SQL"],
            "experiences": [
                {
                    "title": "AI Product Engineer",
                    "company": "Apply Wise",
                    "description": "Built truthful application workflows.",
                }
            ],
            "education": [{"institution": "State University", "degree": "Bachelor of Engineering", "year": "2021"}],
            "certifications": [{"name": "AWS Certified Cloud Practitioner", "issuer": "AWS", "year": "2022"}],
            "projects": [
                {
                    "name": "Resume Parser",
                    "description": "Built a Python parser for resume evidence review.",
                    "technologies": ["Python"],
                }
            ],
        },
    )
    assert edited.status_code == 200
    assert edited.json()["parse_status"] == "reviewed"
    assert edited.json()["skills"] == ["FastAPI", "Python", "SQL"]
    assert edited.json()["experiences"][0]["company"] == "Apply Wise"
    assert edited.json()["projects"][0]["name"] == "Resume Parser"

    reparsed = client.post("/api/v1/resumes/reparse", headers=headers)
    assert reparsed.status_code == 200
    assert reparsed.json()["parse_status"] == "review_required"
    assert reparsed.json()["experiences"][0]["company"] == "Northstar Labs"
    assert reparsed.json()["projects"][0]["name"] == "Apply Wise Copilot"

    activity = client.get("/api/v1/activity/logs", headers=headers).json()
    event_types = {event["event_type"] for event in activity}
    assert {"resume_profile_updated", "resume_reparsed"}.issubset(event_types)


def test_pdf_resume_upload_extracts_profile() -> None:
    token = register_user()["tokens"]["access_token"]
    upload = client.post(
        "/api/v1/resumes/upload",
        headers={"Authorization": f"Bearer {token}"},
        files={"file": ("resume.pdf", make_resume_pdf(), "application/pdf")},
    )
    assert upload.status_code == 201
    assert upload.json()["original_filename"] == "resume.pdf"
    assert upload.json()["skills"] == ["Docker", "Python", "React"]


def test_resume_upload_rejects_unsupported_file() -> None:
    token = register_user()["tokens"]["access_token"]
    upload = client.post(
        "/api/v1/resumes/upload",
        headers={"Authorization": f"Bearer {token}"},
        files={"file": ("resume.txt", b"Summary Skills Python React Docker", "text/plain")},
    )
    assert upload.status_code == 422
    assert upload.json()["detail"] == "Upload a PDF or DOCX resume."


def test_jobs_require_resume_and_saved_job_appears_in_tracker() -> None:
    token = register_user()["tokens"]["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    assert client.get("/api/v1/jobs", headers=headers).json() == []
    upload_resume(token)
    jobs = client.get("/api/v1/jobs", headers=headers).json()
    assert jobs
    saved = client.post(f"/api/v1/jobs/{jobs[0]['id']}/save", headers=headers)
    assert saved.status_code == 201
    assert saved.json()["status"] == "saved"
    refreshed_jobs = client.get("/api/v1/jobs", headers=headers).json()
    saved_job = next(job for job in refreshed_jobs if job["id"] == jobs[0]["id"])
    assert saved_job["saved"] is True
    assert saved_job["application_id"] == saved.json()["id"]
    applications = client.get("/api/v1/applications", headers=headers)
    assert applications.status_code == 200
    assert applications.json() == [saved.json()]


def test_manual_job_capture_is_scored_and_can_be_saved() -> None:
    token = register_user()["tokens"]["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    upload_resume(token)
    captured = client.post(
        "/api/v1/jobs/manual",
        headers=headers,
        json={
            "title": "Platform Engineer",
            "company": "Example Career Site",
            "location": "Remote",
            "url": "https://example.com/careers/platform-engineer",
            "description": "Build reliable Python and FastAPI services with PostgreSQL and Docker for customer-facing applications.",
        },
    )
    assert captured.status_code == 201
    assert captured.json()["source"] == "manual"
    assert captured.json()["match_score"] > 0
    assert "python" in captured.json()["matching_skills"]
    saved = client.post(f"/api/v1/jobs/{captured.json()['id']}/save", headers=headers)
    assert saved.status_code == 201
    captured_again = client.post(
        "/api/v1/jobs/manual",
        headers=headers,
        json={
            "title": "Manual Python Engineer",
            "company": "Example",
            "location": "Remote",
            "url": "https://example.com/careers/platform-engineer",
            "description": "Build reliable Python and FastAPI services with PostgreSQL and Docker for customer-facing applications.",
        },
    )
    assert captured_again.status_code == 201
    assert captured_again.json()["saved"] is True
    assert captured_again.json()["application_id"] == saved.json()["id"]
    approval = client.patch(f"/api/v1/jobs/{captured.json()['id']}/approval", headers=headers, json={"approval_status": "approved"})
    assert approval.status_code == 200
    assert approval.json()["approval_status"] == "approved"
    refreshed_jobs = client.get("/api/v1/jobs?query=Platform%20Engineer&location=Remote", headers=headers)
    assert refreshed_jobs.status_code == 200
    refreshed_capture = next(job for job in refreshed_jobs.json() if job["id"] == captured.json()["id"])
    assert refreshed_capture["title"] == "Platform Engineer"
    assert refreshed_capture["saved"] is True
    assert refreshed_capture["application_id"] == saved.json()["id"]
    assert refreshed_capture["approval_status"] == "approved"


def test_captured_job_can_generate_grounded_resume_and_cover_letter() -> None:
    token = register_user("job-prep@example.com")["tokens"]["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    upload_resume(token)
    captured = client.post(
        "/api/v1/jobs/manual",
        headers=headers,
        json={
            "title": "Platform Engineer",
            "company": "Example Career Site",
            "location": "Remote",
            "url": "https://example.com/careers/platform-engineer",
            "description": "Build reliable Python and FastAPI services with PostgreSQL, Docker, and Kubernetes for customer-facing applications.",
        },
    )
    assert captured.status_code == 201
    saved = client.post(f"/api/v1/jobs/{captured.json()['id']}/save", headers=headers)
    assert saved.status_code == 201

    resume = client.post(f"/api/v1/jobs/{captured.json()['id']}/resume-versions", headers=headers)
    assert resume.status_code == 201
    resume_body = resume.json()
    assert resume_body["job_id"] == captured.json()["id"]
    assert resume_body["prompt_version"] == "applywise-tailoring-v1"
    assert resume_body["model_provider"] in {"deterministic", "gemini"}
    assert resume_body["ats_score"] == captured.json()["ats_score"]
    assert "tailored_bullets" in resume_body["resume_content"]
    generated_text = str(resume_body["resume_content"]).lower()
    assert "python" in generated_text
    assert "kubernetes" not in generated_text
    assert resume_body["missing_skills_analysis"]
    assert resume_body["quality_warnings"] == []

    stored_resume = client.get(f"/api/v1/resume-versions/{resume_body['id']}", headers=headers)
    assert stored_resume.status_code == 200
    assert stored_resume.json()["resume_content"] == resume_body["resume_content"]

    cover_letter = client.post(f"/api/v1/jobs/{captured.json()['id']}/cover-letters", headers=headers)
    assert cover_letter.status_code == 201
    letter_body = cover_letter.json()
    assert letter_body["job_id"] == captured.json()["id"]
    assert letter_body["prompt_version"] == "applywise-tailoring-v1"
    assert "Northstar Labs" in letter_body["content"]
    assert "Kubernetes" not in letter_body["content"]
    assert letter_body["quality_warnings"] == []
    stored_letter = client.get(f"/api/v1/cover-letters/{letter_body['id']}", headers=headers)
    assert stored_letter.status_code == 200
    assert stored_letter.json()["content"] == letter_body["content"]

    review = client.get(f"/api/v1/applications/{saved.json()['id']}/review-before-apply", headers=headers)
    assert review.status_code == 200
    assert review.json()["resume_version_id"] == resume_body["id"]
    assert review.json()["cover_letter_id"] == letter_body["id"]


def test_job_resume_generation_fails_when_resume_is_missing() -> None:
    token = register_user("job-prep-no-resume@example.com")["tokens"]["access_token"]
    user_id = client.get("/api/v1/auth/me", headers={"Authorization": f"Bearer {token}"}).json()["id"]
    job_id = seed_scored_job(user_id)
    response = client.post(f"/api/v1/jobs/{job_id}/resume-versions", headers={"Authorization": f"Bearer {token}"})
    assert response.status_code == 409
    assert "Upload a master resume" in response.json()["detail"]


def test_job_material_generation_is_owner_scoped() -> None:
    owner_token = register_user("job-prep-owner@example.com")["tokens"]["access_token"]
    outsider_token = register_user("job-prep-outsider@example.com")["tokens"]["access_token"]
    owner_headers = {"Authorization": f"Bearer {owner_token}"}
    outsider_headers = {"Authorization": f"Bearer {outsider_token}"}
    upload_resume(owner_token)
    upload_resume(outsider_token)
    captured = client.post(
        "/api/v1/jobs/manual",
        headers=owner_headers,
        json={
            "title": "Owner Scoped Engineer",
            "company": "Private Example",
            "location": "Remote",
            "url": "https://example.com/careers/owner-scoped-engineer",
            "description": "Build Python and FastAPI services with PostgreSQL and Docker.",
        },
    )
    assert captured.status_code == 201

    resume = client.post(f"/api/v1/jobs/{captured.json()['id']}/resume-versions", headers=outsider_headers)
    assert resume.status_code == 404
    cover_letter = client.post(f"/api/v1/jobs/{captured.json()['id']}/cover-letters", headers=outsider_headers)
    assert cover_letter.status_code == 404


def test_job_url_import_fetches_scores_and_persists(monkeypatch: pytest.MonkeyPatch) -> None:
    token = register_user("url-import@example.com")["tokens"]["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    upload_resume(token)

    class FakeAsyncClient:
        def __init__(self, *args: object, **kwargs: object) -> None:
            pass

        async def __aenter__(self) -> "FakeAsyncClient":
            return self

        async def __aexit__(self, *args: object) -> None:
            return None

        async def get(self, url: str, headers: dict[str, str]) -> httpx.Response:
            html = """
            <html>
              <head>
                <title>Platform Engineer - Example Careers</title>
                <meta property="og:site_name" content="Example Careers">
                <meta property="og:title" content="Platform Engineer - Example Careers">
              </head>
              <body>
                <main>
                  <h1>Platform Engineer</h1>
                  <p>Location: Remote India</p>
                  <h2>Job Description</h2>
                  <p>Build reliable Python and FastAPI services with PostgreSQL and Docker for customer-facing applications.</p>
                  <p>Partner with product teams, improve observability, and operate production systems.</p>
                </main>
              </body>
            </html>
            """
            return httpx.Response(200, headers={"content-type": "text/html; charset=utf-8"}, text=html, request=httpx.Request("GET", url))

    monkeypatch.setattr(httpx, "AsyncClient", FakeAsyncClient)
    imported = client.post(
        "/api/v1/jobs/import-url",
        headers=headers,
        json={"url": "https://careers.example.com/platform-engineer"},
    )
    assert imported.status_code == 201
    assert imported.json()["source"] == "url_import"
    assert imported.json()["title"] == "Platform Engineer"
    assert imported.json()["company"] == "Example Careers"
    assert imported.json()["location"] == "Remote India"
    assert imported.json()["match_score"] > 0
    assert "python" in imported.json()["matching_skills"]


def test_job_url_import_rejects_private_urls() -> None:
    token = register_user("url-import-private@example.com")["tokens"]["access_token"]
    upload_resume(token)
    response = client.post(
        "/api/v1/jobs/import-url",
        headers={"Authorization": f"Bearer {token}"},
        json={"url": "http://127.0.0.1:8000/internal-job"},
    )
    assert response.status_code == 422
    assert "public job URLs" in response.text


def test_learning_skill_recommendations_rank_missing_skills() -> None:
    token = register_user("learning@example.com")["tokens"]["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    upload_resume(token)
    captured = client.post(
        "/api/v1/jobs/manual",
        headers=headers,
        json={
            "title": "Platform Engineer",
            "company": "Example",
            "location": "Remote",
            "url": "https://example.com/platform-engineer",
            "description": "Build Kubernetes and Java services with GraphQL for a platform team.",
        },
    )
    assert captured.status_code == 201
    assert client.patch(f"/api/v1/jobs/{captured.json()['id']}/approval", headers=headers, json={"approval_status": "approved"}).status_code == 200

    recommendations = client.get("/api/v1/learning/skill-recommendations", headers=headers)
    assert recommendations.status_code == 200
    skills = {item["skill"]: item for item in recommendations.json()}
    assert skills["java"]["mention_count"] == 1
    assert skills["java"]["approved_role_count"] == 1
    assert "approved" in skills["java"]["reason"]


def test_extension_pairing_code_is_one_time_and_token_is_capture_scoped() -> None:
    token = register_user("extension-owner@example.com")["tokens"]["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    upload_resume(token)
    pairing = client.post("/api/v1/extension/pairing-codes", headers=headers)
    assert pairing.status_code == 201
    code = pairing.json()["code"]

    exchanged = client.post("/api/v1/extension/pair", json={"code": code})
    assert exchanged.status_code == 200
    extension_headers = {"Authorization": f"Bearer {exchanged.json()['extension_token']}"}
    captured = client.post(
        "/api/v1/extension/jobs/capture",
        headers=extension_headers,
        json={
            "title": "Extension Captured Engineer",
            "company": "Example",
            "location": "Remote",
            "url": "https://example.com/careers/extension-captured-engineer",
            "description": "Build Python and FastAPI services with PostgreSQL and Docker for a customer-facing platform.",
        },
    )
    assert captured.status_code == 201
    assert captured.json()["source"] == "manual"
    assert captured.json()["application_id"]
    unreviewed_plan = client.get(f"/api/v1/extension/applications/{captured.json()['application_id']}/autofill-plan", headers=extension_headers)
    assert unreviewed_plan.status_code == 409
    assert "confirm manual review" in unreviewed_plan.json()["detail"]
    client.post(f"/api/v1/applications/{captured.json()['application_id']}/review-before-apply/confirm", headers=headers)
    plan = client.get(f"/api/v1/extension/applications/{captured.json()['application_id']}/autofill-plan", headers=extension_headers)
    assert plan.status_code == 200
    fields = {field["field_key"]: field for field in plan.json()["fields"]}
    assert fields["email"]["value"] == "extension-owner@example.com"
    assert fields["email"]["requires_review"] is True
    assert "first_name" in fields
    assert "last_name" in fields
    assert fields["skills"]["confidence"] >= 0.75
    assert fields["work_authorization"]["requires_review"] is True
    assert fields["work_authorization"]["sensitive"] is True
    assert fields["pronouns"]["requires_review"] is True
    assert fields["pronouns"]["sensitive"] is True
    approved = client.post(
        f"/api/v1/extension/applications/{captured.json()['application_id']}/events",
        headers=extension_headers,
        json={"event_type": "autofill_approved", "metadata": {"approved_field_keys": ["email", "skills"], "approved_field_count": 2}},
    )
    assert approved.status_code == 201
    event = client.post(
        f"/api/v1/extension/applications/{captured.json()['application_id']}/events",
        headers=extension_headers,
        json={"event_type": "autofill_attempted", "metadata": {"filled": [{"field_key": "skills"}], "skipped": []}},
    )
    assert event.status_code == 201
    upload_failed = client.post(
        f"/api/v1/extension/applications/{captured.json()['application_id']}/events",
        headers=extension_headers,
        json={"event_type": "resume_upload_failed", "metadata": {"reason": "browser_blocked_file_path", "manual_upload_required": True}},
    )
    assert upload_failed.status_code == 201
    inspection_failed = client.post(
        f"/api/v1/extension/applications/{captured.json()['application_id']}/events",
        headers=extension_headers,
        json={"event_type": "form_inspection_failed", "metadata": {"reason": "chrome scripting blocked"}},
    )
    assert inspection_failed.status_code == 201
    autofill_failed = client.post(
        f"/api/v1/extension/applications/{captured.json()['application_id']}/events",
        headers=extension_headers,
        json={"event_type": "autofill_failed", "metadata": {"reason": "isolated world blocked"}},
    )
    assert autofill_failed.status_code == 201
    prepared = client.post(f"/api/v1/applications/{captured.json()['application_id']}/prepare", headers=headers)
    assert prepared.status_code == 201
    prepared_events = client.get(f"/api/v1/applications/{captured.json()['application_id']}/events", headers=headers).json()
    assert {"form_inspection_failed", "autofill_failed"}.issubset({event["event_type"] for event in prepared_events})
    assert any(event["event_type"] == "resume_generated" for event in prepared_events)
    assert any(event["event_type"] == "cover_letter_generated" for event in prepared_events)
    downloaded = client.get(f"/api/v1/extension/applications/{captured.json()['application_id']}/resume-download?format=pdf", headers=extension_headers)
    assert downloaded.status_code == 200
    assert downloaded.headers["content-type"] == "application/pdf"
    assert client.get("/api/v1/applications", headers=extension_headers).status_code == 401
    assert client.post("/api/v1/extension/pair", json={"code": code}).status_code == 401


@pytest.mark.parametrize(
    ("payload_patch", "status_code", "error_code", "details_key", "details_value"),
    [
        ({"title": ""}, 422, "missing_title", "missing_fields", ["title"]),
        ({"company": ""}, 422, "missing_company", "missing_fields", ["company"]),
        ({"url": ""}, 422, "missing_url", "missing_fields", ["url"]),
        ({"description": "Too short"}, 422, "short_description", "minimum_length", 40),
    ],
)
def test_extension_capture_returns_structured_validation_errors(
    payload_patch: dict[str, str],
    status_code: int,
    error_code: str,
    details_key: str,
    details_value: object,
) -> None:
    token = register_user(f"extension-errors-{error_code}@example.com")["tokens"]["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    upload_resume(token)
    pairing = client.post("/api/v1/extension/pairing-codes", headers=headers)
    extension_token = client.post("/api/v1/extension/pair", json={"code": pairing.json()["code"]}).json()["extension_token"]
    payload = {
        "title": "LinkedIn Engineer",
        "company": "Example",
        "location": "",
        "url": "https://www.linkedin.com/jobs/view/123/",
        "description": "Build Python services for production users.",
    }
    payload.update(payload_patch)
    response = client.post(
        "/api/v1/extension/jobs/capture",
        headers={"Authorization": f"Bearer {extension_token}"},
        json=payload,
    )
    assert response.status_code == status_code
    assert response.json()["ok"] is False
    assert response.json()["error_code"] == error_code
    assert response.json()["details"][details_key] == details_value


def test_extension_capture_returns_structured_missing_resume_error() -> None:
    token = register_user("extension-no-resume@example.com")["tokens"]["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    pairing = client.post("/api/v1/extension/pairing-codes", headers=headers)
    extension_token = client.post("/api/v1/extension/pair", json={"code": pairing.json()["code"]}).json()["extension_token"]
    response = client.post(
        "/api/v1/extension/jobs/capture",
        headers={"Authorization": f"Bearer {extension_token}"},
        json={
            "title": "LinkedIn Engineer",
            "company": "Example",
            "location": "",
            "url": "https://www.linkedin.com/jobs/view/123/",
            "description": "Build Python services for production users.",
        },
    )
    assert response.status_code == 409
    assert response.json()["error_code"] == "missing_resume"
    assert response.json()["details"]["action"] == "upload_resume"


def test_extension_capture_returns_structured_token_errors() -> None:
    invalid = client.post(
        "/api/v1/extension/jobs/capture",
        headers={"Authorization": "Bearer definitely-not-a-token"},
        json={
            "title": "LinkedIn Engineer",
            "company": "Example",
            "location": "",
            "url": "https://www.linkedin.com/jobs/view/123/",
            "description": "Build Python services for production users.",
        },
    )
    assert invalid.status_code == 401
    assert invalid.json()["error_code"] == "invalid_extension_token"

    settings = get_settings()
    expired_token = jwt.encode(
        {
            "sub": "00000000-0000-0000-0000-000000000000",
            "type": "extension_capture",
            "iss": settings.jwt_issuer,
            "iat": datetime.now(timezone.utc) - timedelta(days=2),
            "exp": datetime.now(timezone.utc) - timedelta(days=1),
        },
        settings.app_secret_key,
        algorithm=settings.jwt_algorithm,
    )
    expired = client.post(
        "/api/v1/extension/jobs/capture",
        headers={"Authorization": f"Bearer {expired_token}"},
        json={
            "title": "LinkedIn Engineer",
            "company": "Example",
            "location": "",
            "url": "https://www.linkedin.com/jobs/view/123/",
            "description": "Build Python services for production users.",
        },
    )
    assert expired.status_code == 401
    assert expired.json()["error_code"] == "expired_extension_token"


def test_extension_capture_returns_structured_database_error(monkeypatch: pytest.MonkeyPatch) -> None:
    token = register_user("extension-db-error@example.com")["tokens"]["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    upload_resume(token)
    pairing = client.post("/api/v1/extension/pairing-codes", headers=headers)
    extension_token = client.post("/api/v1/extension/pair", json={"code": pairing.json()["code"]}).json()["extension_token"]

    async def fail_to_store(*args: object, **kwargs: object) -> object:
        raise SQLAlchemyError("forced db outage")

    monkeypatch.setattr("app.presentation.api_v1.router.score_and_store_manual_job", fail_to_store)
    response = client.post(
        "/api/v1/extension/jobs/capture",
        headers={"Authorization": f"Bearer {extension_token}"},
        json={
            "title": "LinkedIn Engineer",
            "company": "Example",
            "location": "",
            "url": "https://www.linkedin.com/jobs/view/123/",
            "description": "Build Python services for production users.",
        },
    )
    assert response.status_code == 500
    assert response.json()["ok"] is False
    assert response.json()["error_code"] == "database_error"
    assert response.json()["details"]["error"] == "SQLAlchemyError"


def test_linkedin_extension_capture_persists_scores_and_prepares_materials() -> None:
    token = register_user("linkedin-capture-owner@example.com")["tokens"]["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    upload_resume(token)
    pairing = client.post("/api/v1/extension/pairing-codes", headers=headers)
    assert pairing.status_code == 201
    extension_token = client.post("/api/v1/extension/pair", json={"code": pairing.json()["code"]}).json()["extension_token"]
    extension_headers = {"Authorization": f"Bearer {extension_token}"}

    captured = client.post(
        "/api/v1/extension/jobs/capture",
        headers=extension_headers,
        json={
            "title": "AI Senior Software Engineer",
            "company": "Sopra Steria",
            "location": "Noida, Uttar Pradesh, India",
            "url": "https://www.linkedin.com/jobs/view/4425607893/",
            "description": (
                "Company Description About Sopra Steria. Job Description The world is how we shape it. "
                "Role Mission Develop and integrate AI-driven capabilities into the QA Evolution platform. "
                "Ensure secure, auditable, and governed execution of AI agents. Enable scalable agent orchestration "
                "patterns aligned with enterprise architecture. Experience with PydanticAI, LangChain, AutoGen, "
                "OpenAI Agents SDK, Python, Java, Retrieval-Augmented Generation, embeddings, vector databases, "
                "knowledge graphs, LLM evaluation pipelines, and CI/CD integration. Qualifications B Tech in computer science."
            ),
        },
    )

    assert captured.status_code == 201
    body = captured.json()
    assert body["title"] == "AI Senior Software Engineer"
    assert body["company"] == "Sopra Steria"
    assert body["location"] == "Noida, Uttar Pradesh, India"
    assert body["application_id"]
    assert body["ats_score"] is not None
    assert body["match_score"] is not None
    assert body["description"].startswith("Company Description")
    assert "Similar jobs" not in body["description"]

    applications = client.get("/api/v1/applications", headers=headers)
    assert applications.status_code == 200
    saved_application = next(item for item in applications.json() if item["id"] == body["application_id"])
    assert saved_application["title"] == "AI Senior Software Engineer"
    assert saved_application["company"] == "Sopra Steria"

    prepared = client.post(f"/api/v1/applications/{body['application_id']}/prepare", headers=headers)
    assert prepared.status_code == 201
    prepared_body = prepared.json()
    assert prepared_body["resume_version_id"]
    assert prepared_body["cover_letter_id"]
    assert prepared_body["ats_score"] is not None
    assert prepared_body["generation_status"] in {"fallback_no_provider", "generated", "fallback_after_retry"}

    review = client.get(f"/api/v1/applications/{body['application_id']}/review-before-apply", headers=headers)
    assert review.status_code == 200
    assert review.json()["resume_prepared"] is True
    assert review.json()["cover_letter_prepared"] is True
    assert review.json()["manual_final_submission_required"] is True

    events = client.get(f"/api/v1/applications/{body['application_id']}/events", headers=headers).json()
    event_types = {event["event_type"] for event in events}
    assert {"extension_job_captured", "resume_generated", "cover_letter_generated"}.issubset(event_types)


def test_application_status_update_changes_dashboard_summary() -> None:
    token = register_user()["tokens"]["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    upload_resume(token)
    job = client.get("/api/v1/jobs", headers=headers).json()[0]
    application = client.post(f"/api/v1/jobs/{job['id']}/save", headers=headers).json()
    assert client.get("/api/v1/dashboard/summary", headers=headers).json()["saved_jobs"] == 1

    updated = client.patch(f"/api/v1/applications/{application['id']}", headers=headers, json={"status": "interview"})
    assert updated.status_code == 200
    assert updated.json()["status"] == "interview"
    assert updated.json()["applied_at"] is None
    unconfirmed = client.patch(f"/api/v1/applications/{application['id']}", headers=headers, json={"status": "applied"})
    assert unconfirmed.status_code == 409
    assert "manually submitted" in unconfirmed.json()["detail"]
    applied = client.patch(f"/api/v1/applications/{application['id']}", headers=headers, json={"status": "applied", "manual_submission_confirmed": True})
    assert applied.status_code == 200
    assert applied.json()["status"] == "applied"
    assert applied.json()["applied_at"]
    events = client.get(f"/api/v1/applications/{application['id']}/events", headers=headers).json()
    assert any(event["event_type"] == "application_marked_applied" for event in events)
    updated = client.patch(f"/api/v1/applications/{application['id']}", headers=headers, json={"status": "interview"})
    assert updated.status_code == 200
    assert updated.json()["applied_at"] == applied.json()["applied_at"]
    summary = client.get("/api/v1/dashboard/summary", headers=headers).json()
    assert summary["applications"] == 1
    assert summary["interviews"] == 1
    assert summary["saved_jobs"] == 0
    assert sum(summary["weekly_activity"]) == 1
    analytics = client.get("/api/v1/analytics/summary", headers=headers).json()
    assert analytics["total_applications"] == 1
    assert analytics["interview_rate"] == 100


def test_application_update_is_owner_scoped() -> None:
    token = register_user()["tokens"]["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    upload_resume(token)
    job = client.get("/api/v1/jobs", headers=headers).json()[0]
    application = client.post(f"/api/v1/jobs/{job['id']}/save", headers=headers).json()
    outsider = client.post(
        "/api/v1/auth/register",
        json={"email": "outsider@example.com", "password": "correct-horse-battery-staple"},
    ).json()["tokens"]["access_token"]
    response = client.patch(
        f"/api/v1/applications/{application['id']}",
        headers={"Authorization": f"Bearer {outsider}"},
        json={"status": "offer"},
    )
    assert response.status_code == 404


def test_application_notes_create_visible_timeline_event() -> None:
    token = register_user()["tokens"]["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    upload_resume(token)
    job = client.get("/api/v1/jobs", headers=headers).json()[0]
    application = client.post(f"/api/v1/jobs/{job['id']}/save", headers=headers).json()
    updated = client.patch(
        f"/api/v1/applications/{application['id']}",
        headers=headers,
        json={"notes": "Follow up with the recruiter on Friday."},
    )
    assert updated.status_code == 200
    assert updated.json()["notes"] == "Follow up with the recruiter on Friday."
    events = client.get(f"/api/v1/applications/{application['id']}/events", headers=headers)
    assert events.status_code == 200
    assert events.json()[0]["event_type"] == "notes_updated"


def test_interview_round_is_scheduled_and_updates_pipeline() -> None:
    token = register_user()["tokens"]["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    upload_resume(token)
    job = client.get("/api/v1/jobs", headers=headers).json()[0]
    application = client.post(f"/api/v1/jobs/{job['id']}/save", headers=headers).json()
    scheduled = client.post(
        f"/api/v1/applications/{application['id']}/interviews",
        headers=headers,
        json={"interview_round": "Recruiter screen", "scheduled_at": "2026-06-10T10:30:00Z", "notes": "Prepare role summary."},
    )
    assert scheduled.status_code == 201
    assert scheduled.json()["interview_round"] == "Recruiter screen"
    interviews = client.get(f"/api/v1/applications/{application['id']}/interviews", headers=headers)
    assert interviews.json()[0]["notes"] == "Prepare role summary."
    interview_id = interviews.json()[0]["id"]
    updated = client.patch(
        f"/api/v1/applications/{application['id']}/interviews/{interview_id}",
        headers=headers,
        json={"interview_round": "Technical interview", "scheduled_at": "2026-06-12T11:00:00Z", "notes": "Review system design."},
    )
    assert updated.status_code == 200
    assert updated.json()["interview_round"] == "Technical interview"
    assert updated.json()["notes"] == "Review system design."
    removed = client.delete(f"/api/v1/applications/{application['id']}/interviews/{interview_id}", headers=headers)
    assert removed.status_code == 204
    assert client.get(f"/api/v1/applications/{application['id']}/interviews", headers=headers).json() == []
    assert client.get("/api/v1/applications", headers=headers).json()[0]["status"] == "interview"


def test_prepare_application_persists_truthful_resume_and_cover_letter_drafts() -> None:
    token = register_user()["tokens"]["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    upload_resume(token)
    job = client.get("/api/v1/jobs", headers=headers).json()[0]
    application = client.post(f"/api/v1/jobs/{job['id']}/save", headers=headers).json()

    prepared = client.post(f"/api/v1/applications/{application['id']}/prepare", headers=headers)
    assert prepared.status_code == 201
    body = prepared.json()
    assert body["application"]["status"] == "preparing"
    assert body["resume_content"]["draft_type"] == "deterministic_evidence_grounded"
    assert body["resume_content"]["source_resume_id"]
    assert body["resume_content"]["experience"][0]["company"] == "Northstar Labs"
    assert "Northstar Labs" in body["cover_letter"]
    assert body["resume_version_id"]
    assert body["cover_letter_id"]
    assert body["provenance"] == "deterministic_evidence_grounded"
    assert body["generation_status"] == "fallback_no_provider"
    assert body["retry_attempts"] == 1
    versions = client.get("/api/v1/resume-versions", headers=headers)
    assert versions.status_code == 200
    assert versions.json()[0]["id"] == body["resume_version_id"]
    assert versions.json()[0]["job_title"] == job["title"]
    docx = client.get(f"/api/v1/resume-versions/{body['resume_version_id']}/download?format=docx", headers=headers)
    assert docx.status_code == 200
    assert docx.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument")
    assert docx.content.startswith(b"PK")
    pdf = client.get(f"/api/v1/resume-versions/{body['resume_version_id']}/download?format=pdf", headers=headers)
    assert pdf.status_code == 200
    assert pdf.headers["content-type"].startswith("application/pdf")
    assert pdf.content.startswith(b"%PDF")
    export_events = client.get(f"/api/v1/applications/{application['id']}/events", headers=headers)
    assert export_events.status_code == 200
    assert any(event["event_type"] == "resume_export_downloaded" and event["metadata"]["retry_attempts"] == 1 for event in export_events.json())
    version_detail = client.get(f"/api/v1/resume-versions/{body['resume_version_id']}", headers=headers)
    assert version_detail.status_code == 200
    edited_content = dict(version_detail.json()["resume_content"])
    edited_content["professional_summary"] = "Product engineer focused on reliable Python, React, and PostgreSQL systems."
    edited_content["skills"] = ["Python", "React", "PostgreSQL"]
    edited = client.patch(
        f"/api/v1/resume-versions/{body['resume_version_id']}",
        headers=headers,
        json={"resume_content": edited_content},
    )
    assert edited.status_code == 200
    assert edited.json()["resume_content"]["professional_summary"].startswith("Product engineer")
    edited_docx = client.get(f"/api/v1/resume-versions/{body['resume_version_id']}/download?format=docx", headers=headers)
    edited_document = Document(BytesIO(edited_docx.content))
    assert any("Product engineer focused" in paragraph.text for paragraph in edited_document.paragraphs)
    cover_letter = client.get(f"/api/v1/cover-letters/{body['cover_letter_id']}/download", headers=headers)
    assert cover_letter.status_code == 200
    assert cover_letter.headers["content-type"].startswith("text/plain")
    assert "Northstar Labs" in cover_letter.text
    assert "cover-letter-linear-senior-product-engineer.txt" in cover_letter.headers["content-disposition"]
    loaded_cover_letter = client.get(f"/api/v1/cover-letters/{body['cover_letter_id']}", headers=headers)
    assert loaded_cover_letter.status_code == 200
    assert loaded_cover_letter.json()["content"] == cover_letter.text
    edited = client.patch(
        f"/api/v1/cover-letters/{body['cover_letter_id']}",
        headers=headers,
        json={"content": "Dear hiring team,\n\nI am excited to apply with my Python and React experience from Northstar Labs."},
    )
    assert edited.status_code == 200
    assert "excited to apply" in edited.json()["content"]
    downloaded_edit = client.get(f"/api/v1/cover-letters/{body['cover_letter_id']}/download", headers=headers)
    assert "excited to apply" in downloaded_edit.text


def test_assisted_application_checklist_tracks_readiness_without_auto_submit() -> None:
    token = register_user()["tokens"]["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    upload_resume(token)
    job = client.get("/api/v1/jobs", headers=headers).json()[0]
    application = client.post(f"/api/v1/jobs/{job['id']}/save", headers=headers).json()

    checklist_url = f"/api/v1/applications/{application['id']}/assisted-checklist"
    checklist = client.get(checklist_url, headers=headers)
    assert checklist.status_code == 200
    assert checklist.json() == {
        "application_id": application["id"],
        "duplicate_prevention": True,
        "ats_approved": False,
        "resume_prepared": False,
        "resume_version_id": None,
        "cover_letter_prepared": False,
        "cover_letter_id": None,
        "target_url": job["url"],
        "manual_final_submission_required": True,
    }

    approval = client.patch(f"/api/v1/jobs/{job['id']}/approval", headers=headers, json={"approval_status": "approved"})
    assert approval.status_code == 200
    assert client.post(f"/api/v1/applications/{application['id']}/prepare", headers=headers).status_code == 201
    prepared = client.get(checklist_url, headers=headers).json()
    assert prepared["ats_approved"] is True
    assert prepared["resume_prepared"] is True
    assert prepared["resume_version_id"]
    assert prepared["cover_letter_prepared"] is True
    assert prepared["cover_letter_id"]
    assert prepared["manual_final_submission_required"] is True


def test_application_automation_plan_is_dry_run_and_owner_scoped() -> None:
    token = register_user("automation-owner@example.com")["tokens"]["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    upload_resume(token)
    captured = client.post(
        "/api/v1/jobs/manual",
        headers=headers,
        json={
            "title": "LinkedIn Easy Apply Engineer",
            "company": "Example",
            "location": "Remote",
            "url": "https://www.linkedin.com/jobs/view/123",
            "description": "Build Python and FastAPI services with React dashboards and PostgreSQL.",
        },
    )
    saved = client.post(f"/api/v1/jobs/{captured.json()['id']}/save", headers=headers).json()
    plan = client.get(f"/api/v1/applications/{saved['id']}/automation-plan", headers=headers)
    assert plan.status_code == 200
    assert plan.json()["capability"]["site_key"] == "linkedin"
    assert plan.json()["manual_final_submission_required"] is True
    assert plan.json()["steps"][-1]["action"] == "stop_before_submit"
    assert "Prepare a tailored resume" in " ".join(plan.json()["warnings"])

    assert client.post(f"/api/v1/applications/{saved['id']}/prepare", headers=headers).status_code == 201
    prepared_plan = client.get(f"/api/v1/applications/{saved['id']}/automation-plan", headers=headers)
    assert prepared_plan.status_code == 200
    assert "Prepare a tailored resume" not in " ".join(prepared_plan.json()["warnings"])

    outsider = register_user("automation-outsider@example.com")["tokens"]["access_token"]
    assert client.get(f"/api/v1/applications/{saved['id']}/automation-plan", headers={"Authorization": f"Bearer {outsider}"}).status_code == 404


def test_application_automation_runs_are_tracked_and_guarded() -> None:
    token = register_user("automation-run-owner@example.com")["tokens"]["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    upload_resume(token)
    captured = client.post(
        "/api/v1/jobs/manual",
        headers=headers,
        json={
            "title": "Greenhouse Engineer",
            "company": "Example",
            "location": "Remote",
            "url": "https://boards.greenhouse.io/example/jobs/123",
            "description": "Build Python and FastAPI services with React dashboards and PostgreSQL.",
        },
    )
    saved = client.post(f"/api/v1/jobs/{captured.json()['id']}/save", headers=headers).json()
    blocked = client.post(f"/api/v1/applications/{saved['id']}/automation-runs", headers=headers)
    assert blocked.status_code == 201
    assert blocked.json()["status"] == "blocked"
    assert "Prepare a tailored resume" in " ".join(blocked.json()["warnings"])
    assert blocked.json()["manual_final_submission_required"] is True

    runs = client.get(f"/api/v1/applications/{saved['id']}/automation-runs", headers=headers)
    assert runs.status_code == 200
    assert runs.json()[0]["id"] == blocked.json()["id"]

    assert client.post(f"/api/v1/applications/{saved['id']}/prepare", headers=headers).status_code == 201
    ready = client.post(f"/api/v1/applications/{saved['id']}/automation-runs", headers=headers)
    assert ready.status_code == 201
    assert ready.json()["status"] == "ready_for_user_review"
    assert ready.json()["warnings"] == []
    assert ready.json()["result"]["next_action"].startswith("Open the employer page")

    events = client.get(f"/api/v1/applications/{saved['id']}/events", headers=headers).json()
    assert any(event["event_type"] == "automation_run_created" for event in events)
    outsider = register_user("automation-run-outsider@example.com")["tokens"]["access_token"]
    assert client.get(f"/api/v1/applications/{saved['id']}/automation-runs", headers={"Authorization": f"Bearer {outsider}"}).status_code == 404
    assert client.post(f"/api/v1/applications/{saved['id']}/automation-runs", headers={"Authorization": f"Bearer {outsider}"}).status_code == 404


def test_prepare_runs_track_background_material_generation() -> None:
    token = register_user("prepare-run-owner@example.com")["tokens"]["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    upload_resume(token)
    captured = client.post(
        "/api/v1/jobs/manual",
        headers=headers,
        json={
            "title": "Async Prep Engineer",
            "company": "Example",
            "location": "Remote",
            "url": "https://jobs.lever.co/example/123",
            "description": "Build Python and FastAPI services with React dashboards and PostgreSQL.",
        },
    )
    saved = client.post(f"/api/v1/jobs/{captured.json()['id']}/save", headers=headers).json()
    queued = client.post(f"/api/v1/applications/{saved['id']}/prepare-runs", headers=headers)
    assert queued.status_code == 202
    assert queued.json()["job_type"] == "prepare_application_materials"
    assert queued.json()["status"] in {"queued", "running", "succeeded"}

    job = client.get(f"/api/v1/background-jobs/{queued.json()['id']}", headers=headers)
    assert job.status_code == 200
    assert job.json()["status"] == "succeeded"
    assert job.json()["result"]["resume_version_id"]
    assert job.json()["result"]["cover_letter_id"]

    jobs = client.get(f"/api/v1/applications/{saved['id']}/prepare-runs", headers=headers)
    assert jobs.status_code == 200
    assert jobs.json()[0]["id"] == queued.json()["id"]
    events = client.get(f"/api/v1/applications/{saved['id']}/events", headers=headers).json()
    assert any(event["event_type"] == "prepare_job_queued" for event in events)
    assert any(event["event_type"] == "prepare_job_succeeded" for event in events)
    retry_succeeded = client.post(f"/api/v1/background-jobs/{queued.json()['id']}/retry", headers=headers)
    assert retry_succeeded.status_code == 409

    outsider = register_user("prepare-run-outsider@example.com")["tokens"]["access_token"]
    assert client.get(f"/api/v1/background-jobs/{queued.json()['id']}", headers={"Authorization": f"Bearer {outsider}"}).status_code == 404
    assert client.get(f"/api/v1/applications/{saved['id']}/prepare-runs", headers={"Authorization": f"Bearer {outsider}"}).status_code == 404


def test_review_before_apply_packages_materials_and_manual_attention() -> None:
    token = register_user("review-before-apply@example.com")["tokens"]["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    upload_resume(token)
    captured = client.post(
        "/api/v1/jobs/manual",
        headers=headers,
        json={
            "title": "Review Engineer",
            "company": "Example",
            "location": "Remote",
            "url": "https://boards.greenhouse.io/example/jobs/123",
            "description": "Build Python and FastAPI services with React dashboards and PostgreSQL.",
        },
    )
    saved = client.post(f"/api/v1/jobs/{captured.json()['id']}/save", headers=headers).json()
    review = client.get(f"/api/v1/applications/{saved['id']}/review-before-apply", headers=headers)
    assert review.status_code == 200
    assert review.json()["manual_final_submission_required"] is True
    assert review.json()["job_id"] == captured.json()["id"]
    assert review.json()["job_title"] == "Review Engineer"
    assert review.json()["company"] == "Example"
    assert review.json()["match_score"] is not None
    assert review.json()["ats_score"] is not None
    assert review.json()["review_confirmed"] is False
    assert review.json()["resume_prepared"] is False
    assert "Tailored resume" in review.json()["missing_fields"]
    assert "Apply Wise does not submit" in " ".join(review.json()["warnings"])

    assert client.post(f"/api/v1/applications/{saved['id']}/prepare", headers=headers).status_code == 201
    prepared_review = client.get(f"/api/v1/applications/{saved['id']}/review-before-apply", headers=headers)
    assert prepared_review.status_code == 200
    assert prepared_review.json()["resume_prepared"] is True
    assert prepared_review.json()["cover_letter_prepared"] is True
    assert prepared_review.json()["resume_version_id"]
    assert prepared_review.json()["cover_letter_id"]
    assert prepared_review.json()["resume_summary"]
    assert prepared_review.json()["cover_letter_preview"]
    assert "Tailored resume" not in prepared_review.json()["missing_fields"]
    confirmed_review = client.post(f"/api/v1/applications/{saved['id']}/review-before-apply/confirm", headers=headers)
    assert confirmed_review.status_code == 200
    assert confirmed_review.json()["review_confirmed"] is True
    summary = client.get("/api/v1/dashboard/summary", headers=headers).json()
    assert summary["has_resume"] is True
    assert summary["prepared_applications"] == 1
    assert summary["reviewed_applications"] == 1
    reloaded_review = client.get(f"/api/v1/applications/{saved['id']}/review-before-apply", headers=headers)
    assert reloaded_review.json()["review_confirmed"] is True
    version_id = client.get(f"/api/v1/applications/{saved['id']}/assisted-checklist", headers=headers).json()["resume_version_id"]
    version = client.get(f"/api/v1/resume-versions/{version_id}", headers=headers).json()
    edited_content = dict(version["resume_content"])
    edited_content["skills"] = [*edited_content["skills"], "Kubernetes"]
    edited_content["experience"] = [
        *edited_content["experience"],
        {"title": "Principal Engineer", "company": "Acme Corp", "description": "Improved reliability by 45%."},
    ]
    assert client.patch(f"/api/v1/resume-versions/{version_id}", headers=headers, json={"resume_content": edited_content}).status_code == 200
    warned_review = client.get(f"/api/v1/applications/{saved['id']}/review-before-apply", headers=headers)
    assert warned_review.status_code == 200
    assert any("Kubernetes" in warning for warning in warned_review.json()["quality_warnings"])
    assert any("Acme Corp" in warning for warning in warned_review.json()["quality_warnings"])
    blocked_download = client.get(f"/api/v1/resume-versions/{version_id}/download?format=pdf", headers=headers)
    assert blocked_download.status_code == 409
    assert "quality_warnings" in blocked_download.json()["detail"]
    acknowledged_download = client.get(f"/api/v1/resume-versions/{version_id}/download?format=pdf&acknowledge_quality_warnings=true", headers=headers)
    assert acknowledged_download.status_code == 200

    outsider = register_user("review-outsider@example.com")["tokens"]["access_token"]
    assert client.get(f"/api/v1/applications/{saved['id']}/review-before-apply", headers={"Authorization": f"Bearer {outsider}"}).status_code == 404
    assert client.post(f"/api/v1/applications/{saved['id']}/review-before-apply/confirm", headers={"Authorization": f"Bearer {outsider}"}).status_code == 404


def test_resume_download_is_scoped_to_its_owner() -> None:
    owner_token = register_user("owner-download@example.com")["tokens"]["access_token"]
    owner_headers = {"Authorization": f"Bearer {owner_token}"}
    upload_resume(owner_token)
    job = client.get("/api/v1/jobs", headers=owner_headers).json()[0]
    application = client.post(f"/api/v1/jobs/{job['id']}/save", headers=owner_headers).json()
    version_id = client.post(f"/api/v1/applications/{application['id']}/prepare", headers=owner_headers).json()["resume_version_id"]

    other_token = register_user("other-download@example.com")["tokens"]["access_token"]
    response = client.get(
        f"/api/v1/resume-versions/{version_id}/download?format=pdf",
        headers={"Authorization": f"Bearer {other_token}"},
    )
    assert response.status_code == 404
    assert client.get(f"/api/v1/resume-versions/{version_id}", headers={"Authorization": f"Bearer {other_token}"}).status_code == 404
    assert client.patch(
        f"/api/v1/resume-versions/{version_id}",
        headers={"Authorization": f"Bearer {other_token}"},
        json={"resume_content": {"skills": ["Python"], "experience": [{"title": "Other", "company": "Other", "description": "Nope"}]}},
    ).status_code == 404


def test_cover_letter_download_is_scoped_to_its_owner() -> None:
    owner_token = register_user("owner-cover@example.com")["tokens"]["access_token"]
    owner_headers = {"Authorization": f"Bearer {owner_token}"}
    upload_resume(owner_token)
    job = client.get("/api/v1/jobs", headers=owner_headers).json()[0]
    application = client.post(f"/api/v1/jobs/{job['id']}/save", headers=owner_headers).json()
    cover_letter_id = client.post(f"/api/v1/applications/{application['id']}/prepare", headers=owner_headers).json()["cover_letter_id"]

    other_token = register_user("other-cover@example.com")["tokens"]["access_token"]
    response = client.get(
        f"/api/v1/cover-letters/{cover_letter_id}/download",
        headers={"Authorization": f"Bearer {other_token}"},
    )
    assert response.status_code == 404
    assert client.get(f"/api/v1/cover-letters/{cover_letter_id}", headers={"Authorization": f"Bearer {other_token}"}).status_code == 404
    assert client.patch(
        f"/api/v1/cover-letters/{cover_letter_id}",
        headers={"Authorization": f"Bearer {other_token}"},
        json={"content": "This edit should not be allowed for another user's cover letter."},
    ).status_code == 404


def test_register_and_read_current_user() -> None:
    registration = client.post(
        "/api/v1/auth/register",
        json={"email": "Avery@example.com", "password": "correct-horse-battery-staple"},
    )
    assert registration.status_code == 201
    body = registration.json()
    assert body["user"]["email"] == "avery@example.com"
    assert body["user"]["role"] == "member"
    assert body["tokens"]["token_type"] == "bearer"
    assert registration.cookies.get("jobcopilot_refresh")

    current_user = client.get(
        "/api/v1/auth/me",
        headers={"Authorization": f"Bearer {body['tokens']['access_token']}"},
    )
    assert current_user.status_code == 200
    assert current_user.json()["email"] == "avery@example.com"
    assert current_user.json()["role"] == "member"


def test_allowlisted_owner_registers_as_admin() -> None:
    registration = client.post(
        "/api/v1/auth/register",
        json={"email": "OWNER@example.com", "password": "correct-horse-battery-staple"},
    )
    assert registration.status_code == 201
    assert registration.json()["user"]["role"] == "admin"


def test_owner_subscription_bypass_and_member_free_tier() -> None:
    owner = register_user("owner@example.com")
    owner_subscription = client.get(
        "/api/v1/billing/subscription",
        headers={"Authorization": f"Bearer {owner['tokens']['access_token']}"},
    )
    assert owner_subscription.status_code == 200
    assert owner_subscription.json()["plan"] == "owner"
    assert owner_subscription.json()["owner_bypass"] is True
    assert "gemini_tailoring" in owner_subscription.json()["features"]

    member = register_user("member-subscription@example.com")
    member_subscription = client.get(
        "/api/v1/billing/subscription",
        headers={"Authorization": f"Bearer {member['tokens']['access_token']}"},
    )
    assert member_subscription.status_code == 200
    assert member_subscription.json()["plan"] == "free"
    assert member_subscription.json()["owner_bypass"] is False


def test_admin_summary_is_owner_only() -> None:
    owner = register_user("owner@example.com")
    owner_response = client.get(
        "/api/v1/admin/summary",
        headers={"Authorization": f"Bearer {owner['tokens']['access_token']}"},
    )
    assert owner_response.status_code == 200
    assert owner_response.json()["users"] >= 1
    assert "applications" in owner_response.json()

    member = register_user("member-admin-summary@example.com")
    member_response = client.get(
        "/api/v1/admin/summary",
        headers={"Authorization": f"Bearer {member['tokens']['access_token']}"},
    )
    assert member_response.status_code == 403


def test_existing_allowlisted_owner_becomes_admin_on_login(monkeypatch) -> None:
    payload = {"email": "avery@example.com", "password": "correct-horse-battery-staple"}
    assert client.post("/api/v1/auth/register", json=payload).json()["user"]["role"] == "member"

    monkeypatch.setattr(get_settings(), "admin_emails", payload["email"])
    login = client.post("/api/v1/auth/login", json=payload)
    assert login.status_code == 200
    assert login.json()["user"]["role"] == "admin"


def test_current_user_requires_access_token() -> None:
    current_user = client.get("/api/v1/auth/me")
    assert current_user.status_code == 401


def test_duplicate_registration_is_rejected() -> None:
    payload = {"email": "avery@example.com", "password": "correct-horse-battery-staple"}
    assert client.post("/api/v1/auth/register", json=payload).status_code == 201
    assert client.post("/api/v1/auth/register", json=payload).status_code == 409


def test_login_rejects_invalid_password() -> None:
    payload = {"email": "avery@example.com", "password": "correct-horse-battery-staple"}
    assert client.post("/api/v1/auth/register", json=payload).status_code == 201

    login = client.post(
        "/api/v1/auth/login",
        json={"email": payload["email"], "password": "definitely-not-correct"},
    )
    assert login.status_code == 401


def test_authenticated_user_can_change_password_and_revoke_refresh_tokens() -> None:
    payload = {"email": "avery@example.com", "password": "correct-horse-battery-staple"}
    registration = client.post("/api/v1/auth/register", json=payload)
    assert registration.status_code == 201
    access_token = registration.json()["tokens"]["access_token"]
    refresh_token = registration.cookies.get("jobcopilot_refresh")
    assert refresh_token

    wrong_current = client.post(
        "/api/v1/auth/change-password",
        headers={"Authorization": f"Bearer {access_token}"},
        json={"current_password": "wrong-password", "new_password": "new-correct-horse-password"},
    )
    assert wrong_current.status_code == 401

    changed = client.post(
        "/api/v1/auth/change-password",
        headers={"Authorization": f"Bearer {access_token}"},
        json={"current_password": payload["password"], "new_password": "new-correct-horse-password"},
    )
    assert changed.status_code == 204

    assert client.post("/api/v1/auth/login", json=payload).status_code == 401
    assert client.post(
        "/api/v1/auth/login",
        json={"email": payload["email"], "password": "new-correct-horse-password"},
    ).status_code == 200
    client.cookies.set("jobcopilot_refresh", refresh_token, path="/api/v1/auth")
    assert client.post("/api/v1/auth/refresh").status_code == 401


def test_password_reset_changes_password_and_rejects_token_reuse() -> None:
    payload = {"email": "reset@example.com", "password": "correct-horse-battery-staple"}
    assert client.post("/api/v1/auth/register", json=payload).status_code == 201
    requested = client.post("/api/v1/auth/password-reset/request", json={"email": payload["email"]})
    assert requested.status_code == 200
    assert requested.json()["message"].startswith("If an account exists")
    token = requested.json()["reset_url"].split("token=", 1)[1]

    reset = client.post(
        "/api/v1/auth/password-reset/confirm",
        json={"token": token, "new_password": "new-correct-horse-password"},
    )
    assert reset.status_code == 204
    assert client.post("/api/v1/auth/login", json=payload).status_code == 401
    assert client.post(
        "/api/v1/auth/login",
        json={"email": payload["email"], "password": "new-correct-horse-password"},
    ).status_code == 200
    reused = client.post(
        "/api/v1/auth/password-reset/confirm",
        json={"token": token, "new_password": "another-new-password"},
    )
    assert reused.status_code == 401


def test_password_reset_request_has_support_recovery_path_when_reset_url_is_hidden(monkeypatch) -> None:
    settings = get_settings()
    monkeypatch.setattr(settings, "app_env", "production")

    requested = client.post("/api/v1/auth/password-reset/request", json={"email": "support-reset@example.com"})

    assert requested.status_code == 200
    body = requested.json()
    assert body["reset_url"] is None
    assert "support@applywise.app" in body["message"]


def test_password_reset_request_sends_email_when_smtp_is_enabled(monkeypatch) -> None:
    payload = {"email": "smtp-reset@example.com", "password": "correct-horse-battery-staple"}
    assert client.post("/api/v1/auth/register", json=payload).status_code == 201
    settings = get_settings()
    monkeypatch.setattr(settings, "app_env", "production")
    monkeypatch.setattr(settings, "frontend_base_url", "https://apply-wise-1.vercel.app")
    monkeypatch.setattr(settings, "smtp_host", "smtp.example.com")
    monkeypatch.setattr(settings, "smtp_from_email", "support@applywise.app")
    sent: list[tuple[str, str]] = []

    def fake_send_password_reset_email(_settings, recipient: str, reset_url: str) -> None:
        sent.append((recipient, reset_url))

    monkeypatch.setattr("app.presentation.api_v1.auth.send_password_reset_email", fake_send_password_reset_email)
    requested = client.post("/api/v1/auth/password-reset/request", json={"email": payload["email"]})

    assert requested.status_code == 200
    assert requested.json()["reset_url"] is None
    assert "password reset instructions" in requested.json()["message"]
    assert sent and sent[0][0] == payload["email"]
    assert sent[0][1].startswith("https://apply-wise-1.vercel.app/reset-password?token=")


def test_refresh_token_rotation_detects_reuse() -> None:
    registration = client.post(
        "/api/v1/auth/register",
        json={"email": "avery@example.com", "password": "correct-horse-battery-staple"},
    )
    original_token = registration.cookies.get("jobcopilot_refresh")
    assert original_token

    client.cookies.set("jobcopilot_refresh", original_token, path="/api/v1/auth")
    refresh = client.post("/api/v1/auth/refresh")
    assert refresh.status_code == 200
    rotated_token = refresh.cookies.get("jobcopilot_refresh")
    assert rotated_token
    assert rotated_token != original_token

    client.cookies.set("jobcopilot_refresh", original_token, path="/api/v1/auth")
    replay = client.post("/api/v1/auth/refresh")
    assert replay.status_code == 401
    assert replay.json()["detail"] == "Refresh token reuse detected."

    client.cookies.set("jobcopilot_refresh", rotated_token, path="/api/v1/auth")
    revoked_family = client.post("/api/v1/auth/refresh")
    assert revoked_family.status_code == 401


def test_logout_revokes_refresh_token() -> None:
    registration = client.post(
        "/api/v1/auth/register",
        json={"email": "avery@example.com", "password": "correct-horse-battery-staple"},
    )
    refresh_token = registration.cookies.get("jobcopilot_refresh")
    assert refresh_token

    client.cookies.set("jobcopilot_refresh", refresh_token, path="/api/v1/auth")
    logout = client.post("/api/v1/auth/logout")
    assert logout.status_code == 204

    client.cookies.set("jobcopilot_refresh", refresh_token, path="/api/v1/auth")
    refresh = client.post("/api/v1/auth/refresh")
    assert refresh.status_code == 401


def register_user(email: str = "resume-owner@example.com") -> dict:
    response = client.post(
        "/api/v1/auth/register",
        json={"email": email, "password": "correct-horse-battery-staple"},
    )
    assert response.status_code == 201
    return response.json()


def make_resume_docx() -> bytes:
    document = Document()
    document.add_heading("Summary", level=1)
    document.add_paragraph("Product engineer building reliable user-facing software.")
    document.add_heading("Experience", level=1)
    document.add_paragraph("Senior Product Engineer")
    document.add_paragraph("Northstar Labs")
    document.add_paragraph("Built React and TypeScript applications with Python, FastAPI, PostgreSQL, and Docker.")
    document.add_heading("Education", level=1)
    document.add_paragraph("State University - Bachelor of Engineering - 2021")
    document.add_heading("Certifications", level=1)
    document.add_paragraph("AWS Certified Cloud Practitioner - Amazon Web Services - 2022")
    document.add_heading("Projects", level=1)
    document.add_paragraph("Apply Wise Copilot")
    document.add_paragraph("Built a React and FastAPI workflow that parses resumes, scores jobs, and tracks applications.")
    document.add_heading("Skills", level=1)
    document.add_paragraph("React TypeScript Python FastAPI PostgreSQL Docker")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def upload_resume(token: str) -> None:
    response = client.post(
        "/api/v1/resumes/upload",
        headers={"Authorization": f"Bearer {token}"},
        files={"file": ("resume.docx", make_resume_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 201


def seed_scored_job(user_id: str) -> str:
    job_id = uuid4()

    async def seed() -> None:
        async with session_factory() as session:
            session.add(
                Job(
                    id=job_id,
                    source="test",
                    external_id=str(job_id),
                    title="Seeded Platform Engineer",
                    company="Seeded Example",
                    location="Remote",
                    salary=None,
                    url=f"https://example.com/jobs/{job_id}",
                    description="Build Python and FastAPI services with PostgreSQL and Docker.",
                    description_hash=str(job_id),
                )
            )
            session.add(
                JobMatch(
                    user_id=UUID(user_id),
                    job_id=job_id,
                    match_score=78,
                    ats_score=72,
                    ats_breakdown={"keywords": 70, "sections": 80},
                    ats_recommendations=["Add stronger Python service keywords if supported by your resume."],
                    matching_skills=["python", "fastapi", "postgresql", "docker"],
                    missing_skills=["kubernetes"],
                )
            )
            await session.commit()

    asyncio.run(seed())
    return str(job_id)


def make_resume_pdf() -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Summary\nExperience\nSkills\nPython React Docker\nEducation")
    content = document.tobytes()
    document.close()
    return content
