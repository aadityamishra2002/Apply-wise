from dataclasses import dataclass
from io import BytesIO
import re

import fitz
from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph

from app.infrastructure.db.models import Job, ResumeTemplate, ResumeVersion, User, UserProfile


PLACEHOLDER_PATTERN = re.compile(r"\{\{\s*([A-Z_]+)\s*\}\}", re.IGNORECASE)
SUPPORTED_PLACEHOLDERS = {
    "NAME",
    "CONTACT",
    "SUMMARY",
    "PROFESSIONAL_SUMMARY",
    "SKILLS",
    "EXPERIENCE",
    "PROJECTS",
    "EDUCATION",
    "CERTIFICATIONS",
}
SECTION_ORDER = ("SUMMARY", "SKILLS", "EXPERIENCE", "PROJECTS", "EDUCATION", "CERTIFICATIONS")


@dataclass(frozen=True)
class ResumeIdentity:
    name: str
    contact: str


class ResumeExportService:
    def preview_sections(self, content: dict, *, user: User | None = None, user_profile: UserProfile | None = None) -> dict[str, list[str]]:
        identity = build_identity(user, user_profile)
        return render_sections(normalized_resume_content(content), identity)

    def build_docx(
        self,
        version: ResumeVersion,
        job: Job,
        *,
        user: User | None = None,
        user_profile: UserProfile | None = None,
        template: ResumeTemplate | None = None,
    ) -> bytes:
        identity = build_identity(user, user_profile)
        document = load_template_document(template) if template else build_default_template()
        render_template(document, version.resume_content, identity)
        buffer = BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    def build_pdf(
        self,
        version: ResumeVersion,
        job: Job,
        *,
        user: User | None = None,
        user_profile: UserProfile | None = None,
        template: ResumeTemplate | None = None,
    ) -> bytes:
        identity = build_identity(user, user_profile)
        content = normalized_resume_content(version.resume_content)
        section_order = complete_section_order(template)
        lines = build_plain_lines(content, identity, section_order)
        document = fitz.open()
        page = document.new_page()
        y = 42
        for line, size, bold in lines:
            wrapped = wrap_text(line, width=78 if size >= 12 else 92)
            for part in wrapped:
                if y > 780:
                    page = document.new_page()
                    y = 42
                page.insert_text((54, y), part, fontsize=size, fontname="helv" if not bold else "helv")
                y += size + 5
            if not line:
                y += 4
        output = document.tobytes()
        document.close()
        return output


def build_default_template() -> DocxDocument:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    for style_name in ("Heading 1", "Heading 2"):
        style = document.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11 if style_name == "Heading 1" else 10)
        style.font.bold = True
    name = document.add_paragraph("{{NAME}}")
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.style = document.styles["Title"]
    contact = document.add_paragraph("{{CONTACT}}")
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for section_name in SECTION_ORDER:
        document.add_heading(section_name.replace("_", " ").title(), level=1)
        document.add_paragraph(f"{{{{{section_name}}}}}")
    return document


def load_template_document(template: ResumeTemplate) -> DocxDocument:
    return Document(BytesIO(template.template_bytes))


def extract_placeholders_from_docx(content: bytes) -> list[str]:
    document = Document(BytesIO(content))
    found: list[str] = []
    for paragraph in document.paragraphs:
        for match in PLACEHOLDER_PATTERN.findall(paragraph.text):
            name = match.strip().upper()
            if name in SUPPORTED_PLACEHOLDERS and name not in found:
                found.append(name)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for match in PLACEHOLDER_PATTERN.findall(paragraph.text):
                        name = match.strip().upper()
                        if name in SUPPORTED_PLACEHOLDERS and name not in found:
                            found.append(name)
    return found


def render_template(document: DocxDocument, raw_content: dict, identity: ResumeIdentity) -> None:
    content = normalized_resume_content(raw_content)
    sections = render_sections(content, identity)
    replaced: set[str] = set()
    paragraphs = all_paragraphs(document)
    for paragraph in paragraphs:
        placeholders = [match.strip().upper() for match in PLACEHOLDER_PATTERN.findall(paragraph.text)]
        for placeholder in placeholders:
            canonical = "SUMMARY" if placeholder == "PROFESSIONAL_SUMMARY" else placeholder
            if canonical not in sections:
                continue
            replace_placeholder_paragraph(paragraph, sections[canonical])
            replaced.add(canonical)
            break
    if not replaced:
        append_rendered_sections(document, sections)
    else:
        append_missing_sections(document, sections, replaced)


def all_paragraphs(document: DocxDocument) -> list[Paragraph]:
    paragraphs = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    return paragraphs


def replace_placeholder_paragraph(paragraph: Paragraph, lines: list[str]) -> None:
    if not lines:
        clear_paragraph(paragraph)
        return
    original_style = paragraph.style
    current = paragraph
    for index, line in enumerate(lines):
        if index > 0:
            current = insert_paragraph_after(current)
            current.style = original_style
        set_paragraph_text(current, line)


def append_rendered_sections(document: DocxDocument, sections: dict[str, list[str]]) -> None:
    if document.paragraphs:
        document.add_section(WD_SECTION.CONTINUOUS)
    if sections["NAME"]:
        paragraph = document.add_paragraph(sections["NAME"][0])
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if sections["CONTACT"]:
        paragraph = document.add_paragraph(sections["CONTACT"][0])
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for section_name in SECTION_ORDER:
        lines = sections.get(section_name, [])
        if not lines:
            continue
        document.add_heading(section_name.replace("_", " ").title(), level=1)
        for line in lines:
            document.add_paragraph(line)


def append_missing_sections(document: DocxDocument, sections: dict[str, list[str]], replaced: set[str]) -> None:
    missing_sections = [section_name for section_name in SECTION_ORDER if section_name not in replaced and sections.get(section_name)]
    if not missing_sections:
        return
    document.add_section(WD_SECTION.CONTINUOUS)
    for section_name in missing_sections:
        document.add_heading(section_name.replace("_", " ").title(), level=1)
        for line in sections[section_name]:
            document.add_paragraph(line)


def insert_paragraph_after(paragraph: Paragraph) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        first_run = paragraph.runs[0]
        for run in paragraph.runs[1:]:
            run.text = ""
        first_run.text = text
    else:
        paragraph.add_run(text)


def clear_paragraph(paragraph: Paragraph) -> None:
    for run in paragraph.runs:
        run.text = ""


def render_sections(content: dict, identity: ResumeIdentity) -> dict[str, list[str]]:
    return {
        "NAME": [identity.name] if identity.name else [],
        "CONTACT": [identity.contact] if identity.contact else [],
        "SUMMARY": [content["professional_summary"]] if content["professional_summary"] else [],
        "SKILLS": [", ".join(content["skills"])] if content["skills"] else [],
        "EXPERIENCE": render_experience(content["experience"]),
        "PROJECTS": render_projects(content["projects"]),
        "EDUCATION": render_education(content["education"]),
        "CERTIFICATIONS": render_certifications(content["certifications"]),
    }


def normalized_resume_content(content: dict) -> dict:
    return {
        "professional_summary": str(content.get("professional_summary", "")).strip(),
        "skills": [str(skill).strip() for skill in content.get("skills", []) if str(skill).strip()],
        "experience": [item for item in content.get("experience", []) if isinstance(item, dict)],
        "projects": [item for item in content.get("projects", []) if isinstance(item, dict)],
        "education": [item for item in content.get("education", []) if isinstance(item, dict)],
        "certifications": [item for item in content.get("certifications", []) if isinstance(item, dict)],
    }


def render_experience(items: list[dict]) -> list[str]:
    lines: list[str] = []
    for item in items:
        heading = join_non_empty([item.get("title"), item.get("company")], " | ")
        description = str(item.get("description", "")).strip()
        if heading:
            lines.append(heading)
        if description:
            lines.append(description)
    return lines


def render_projects(items: list[dict]) -> list[str]:
    lines: list[str] = []
    for item in items:
        technologies = item.get("technologies") if isinstance(item.get("technologies"), list) else []
        heading = str(item.get("name", "")).strip()
        description = str(item.get("description", "")).strip()
        tech_line = ", ".join(str(tech).strip() for tech in technologies if str(tech).strip())
        if heading:
            lines.append(heading)
        if description:
            lines.append(description)
        if tech_line:
            lines.append(f"Technologies: {tech_line}")
    return lines


def render_education(items: list[dict]) -> list[str]:
    lines: list[str] = []
    for item in items:
        value = join_non_empty([item.get("degree"), item.get("institution"), item.get("year")], " | ")
        if value:
            lines.append(value)
    return lines


def render_certifications(items: list[dict]) -> list[str]:
    lines: list[str] = []
    for item in items:
        value = join_non_empty([item.get("name"), item.get("issuer"), item.get("year")], " | ")
        if value:
            lines.append(value)
    return lines


def build_plain_lines(content: dict, identity: ResumeIdentity, section_order: list[str]) -> list[tuple[str, int, bool]]:
    sections = render_sections(content, identity)
    lines: list[tuple[str, int, bool]] = []
    if sections["NAME"]:
        lines.append((sections["NAME"][0], 16, True))
    if sections["CONTACT"]:
        lines.append((sections["CONTACT"][0], 9, False))
    lines.append(("", 7, False))
    for section_name in section_order:
        values = sections.get(section_name, [])
        if not values:
            continue
        lines.append((section_name.replace("_", " "), 11, True))
        for value in values:
            lines.append((value, 9, False))
        lines.append(("", 7, False))
    return lines


def build_identity(user: User | None, user_profile: UserProfile | None) -> ResumeIdentity:
    name = (user_profile.full_name if user_profile else "").strip()
    if not name and user:
        name = user.email.split("@")[0].replace(".", " ").replace("_", " ").title()
    contact_parts = []
    if user:
        contact_parts.append(user.email)
    if user_profile:
        contact_parts.extend([user_profile.phone, user_profile.location, user_profile.linkedin_url, user_profile.github_url, user_profile.portfolio_url])
    return ResumeIdentity(name=name, contact=" | ".join(part.strip() for part in contact_parts if part and part.strip()))


def placeholder_section_order(template: ResumeTemplate | None) -> list[str]:
    if not template:
        return []
    order: list[str] = []
    for placeholder in template.placeholders or []:
        normalized = "SUMMARY" if str(placeholder).upper() == "PROFESSIONAL_SUMMARY" else str(placeholder).upper()
        if normalized in SECTION_ORDER and normalized not in order:
            order.append(normalized)
    return order


def complete_section_order(template: ResumeTemplate | None) -> list[str]:
    order = placeholder_section_order(template)
    if not order:
        return list(SECTION_ORDER)
    return order + [section for section in SECTION_ORDER if section not in order]


def join_non_empty(values: list[object], separator: str) -> str:
    return separator.join(str(value).strip() for value in values if value is not None and str(value).strip())


def wrap_text(value: str, width: int = 92) -> list[str]:
    words = value.split()
    if not words:
        return [""]
    lines: list[str] = []
    current = words[0]
    for word in words[1:]:
        if len(current) + len(word) + 1 <= width:
            current = f"{current} {word}"
        else:
            lines.append(current)
            current = word
    lines.append(current)
    return lines


def download_filename(job: Job, extension: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", f"{job.company}-{job.title}".lower()).strip("-")
    return f"{slug or 'tailored-resume'}.{extension}"
